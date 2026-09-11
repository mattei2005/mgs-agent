from openpyxl import load_workbook
import pathlib,json,decimal,collections,hashlib
D=decimal.Decimal;root=pathlib.Path('/root/mgs-agent/work/adops-revenue-1547811852405178460');source=pathlib.Path('/root/.hermes/profiles/zeus/cache/documents/doc_10d5b4efd431_report_1_de_set_ate_9_de_set.xlsx');rows=json.loads((root/'normalized-rows.json').read_text());old=collections.defaultdict(D);counts=collections.Counter();w=load_workbook(source,read_only=True,data_only=True)
for s in w:
 currency=s.title.upper()
 for r in s.iter_rows(min_row=2,values_only=True):
  if all(v is None for v in r):continue
  placement=str(r[1]).strip();medium=str(r[2]).strip();value=D(str(r[9]));old[(currency,placement,medium)]+=value;counts[currency]+=1
w.close();new=collections.defaultdict(D)
for r in rows:
 if r['month']!='2026-09':continue
 new[(r['currency'],r['placement'],r['medium'])]+=D(r['revenue'])
keys=set(old)|set(new);diff=[]
for k in sorted(keys):
 d=new[k]-old[k]
 if abs(d)>D('0.000000001'):diff.append({'currency':k[0],'placement':k[1],'medium':k[2],'old':str(old[k]),'new':str(new[k]),'difference':str(d)})
byplacement=collections.defaultdict(lambda:{'old':D(0),'new':D(0)})
for (c,p,m),v in old.items():byplacement[(c,p)]['old']+=v
for (c,p,m),v in new.items():byplacement[(c,p)]['new']+=v
pdiff=[{'currency':c,'placement':p,'old':str(v['old']),'new':str(v['new']),'difference':str(v['new']-v['old'])} for (c,p),v in sorted(byplacement.items()) if abs(v['new']-v['old'])>D('0.000000001')];out={'pass':True,'old_source':str(source),'old_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'old_rows':dict(counts),'old_totals':{c:str(sum((v for (cc,_,_),v in old.items() if cc==c),D(0))) for c in ['USD','CAD']},'new_totals':{c:str(sum((v for (cc,_,_),v in new.items() if cc==c),D(0))) for c in ['USD','CAD']},'group_differences':diff,'placement_differences':pdiff,'all_usd_match':not any(x['currency']=='USD' for x in diff),'cad_difference_count':sum(x['currency']=='CAD' for x in diff)};(root/'compare-september-original.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n');print(json.dumps({'pass':True,'totals':{'old':out['old_totals'],'new':out['new_totals']},'placement_differences':pdiff,'group_difference_count':len(diff),'cad_difference_count':out['cad_difference_count'],'all_usd_match':out['all_usd_match']},ensure_ascii=False))
