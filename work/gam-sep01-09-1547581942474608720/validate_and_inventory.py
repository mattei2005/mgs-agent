import json,os,fcntl,hashlib
from pathlib import Path
from datetime import datetime,timezone
from decimal import Decimal as D
from collections import defaultdict,Counter
from openpyxl import load_workbook
P=Path(__file__).parent
summary=json.loads((P/'summary.json').read_text());source=Path(summary['input']);output=Path(summary['output'])
assert hashlib.sha256(source.read_bytes()).hexdigest()==summary['source_sha256']
assert hashlib.sha256(output.read_bytes()).hexdigest()==summary['output_sha256']
lineage=json.loads((P/'lineage.json').read_text());index={(r['sheet'],r['row']):r for r in lineage};assert len(index)==27091
w=load_workbook(source,read_only=True,data_only=True);count=0;raw=defaultdict(lambda:D(0))
for s in w:
 for n,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
  if all(v is None for v in row):continue
  r=index[(s.title,n)];assert r['amount']==str(D(str(row[9])))
  assert r['placement']==row[1] and r['medium']==row[2] and r['campaign']==row[3]
  assert r['key'][1]==row[0].date().isoformat();count+=1;raw[(s.title.upper(),row[0].date().isoformat())]+=D(str(row[9]))
assert count==len(index)
# Independent rules regression against every matched source line.
for r in lineage:
 k=r['key'];p=r['placement']
 if p=='pl_digital-trust_eggbev_gb':assert k[3]=='gb-cc-en'
 if p=='pl_digital-trust_eggbev_us':assert k[3]=='us-cc-en'
 if p=='pl_digital-trust_openzedfinanzas_es':assert k[3]=='es-cc-es'
 if p=='pl_digital-trust_openzedfinanzas_us':assert k[3]=='us-cc-es'
 if p.startswith('pl_digital-trust_topfeedfinanzas_'):assert k[3]=='us-cc-es'
 if p.startswith('pl_digital-trust_gamezonead_'):assert k[3:]==['br-game-br','g002-s']
 if p.startswith('pl_digital-trust_autocreditadx_'):assert k[3]=='us-car-en'
 if p=='pl_digital-trust_yolokfx_us' and 'Residual G002 autorizado, sem identidade segura' in r['rules']:assert k[4]=='g002-s'
 if p=='pl_digital-trust_yolokfx_us' and 'Medium GAM explícito' in r['rules']:assert k[4]==r['medium']
wb=load_workbook(output,data_only=False);assert wb.sheetnames==['Receita USD','Receita CAD','Validacao']
assert not any(c.data_type in ('f','e') for s in wb for row in s for c in row)
seen=Counter()
for currency in ['USD','CAD']:
 s=wb['Receita '+currency];expected=summary['groups_by_currency'][currency];assert s.max_row==expected+2
 assert s.auto_filter.ref==f'A1:E{s.max_row-1}'
 for row in s.iter_rows(min_row=2,max_row=s.max_row-1,values_only=True):seen[(currency,row[0])]+=1
assert set(seen)==set(raw)
verification={'status':'PASS','original_rows_readback':count,'source_lineage_matches':True,'approved_rule_regressions':True,'currency_days':len(raw),'output_groups':summary['output_groups'],'source_hash_unchanged':True,'formula_or_error_cells':0,'classifications_pending':len(summary['pending'])}
(P/'independent-verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2))
root=Path('/root/mgs-agent');ident='zeus-gam-sep01-09-1547581942474608720';now=datetime.now(timezone.utc).isoformat()
paths=[p for p in P.iterdir() if p.is_file()]+[root/'data/agent-checkpoints.json',Path('/root/.hermes/profiles/zeus/skills/ops/smartbidding-dashboard-map/references/adgroup-gam-revenue-reconciliation.md')]
entry={'id':ident,'agent':'zeus','type':'finance_excel_consolidation','authorization_message_id':'1547581942474608720','source_thread_id':'1545426987756298340','updated_at':now,'evidence_directory':str(P),'paths':[{'path':str(p),'size_bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(paths)],'summary':verification,'financial_sheet_writes':False,'report_infra_pending':True}
p=root/'data/infra-inventory.json'
with (root/'data/.infra-inventory.lock').open('a') as f:
 fcntl.flock(f,fcntl.LOCK_EX);d=json.loads(p.read_text());rows=d.setdefault('runtime_artifacts',[]);rows[:]=[x for x in rows if x.get('id')!=ident];rows.append(entry)
 tmp=p.with_name(p.name+'.'+ident+'.tmp')
 with tmp.open('w') as out:json.dump(d,out,ensure_ascii=False,indent=2);out.write('\n');out.flush();os.fsync(out.fileno())
 tmp.chmod(p.stat().st_mode&0o777);os.replace(tmp,p);check=json.loads(p.read_text());assert next(x for x in check['runtime_artifacts'] if x.get('id')==ident)==entry
with (root/'logs/events-audit.jsonl').open('a') as f:
 fcntl.flock(f,fcntl.LOCK_EX);f.write(json.dumps({'timestamp':now,'agent':'zeus','event':'gam_september_1_9_excel_validated','authorization_message_id':'1547581942474608720','artifact_id':ident,'evidence':verification,'totals':summary['raw_totals'],'pending_cad':summary['pending_total_cad'],'financial_sheet_writes':False,'recovered_tool_errors':['python alias absent: used python3','checkpoint subcommand corrected to checkpoint-upsert','read-only POST201 accepted after payload validation','read-only shell source inspections blocked by gateway guard; no restart attempted'],'source':str(P)},ensure_ascii=False)+'\n');f.flush();os.fsync(f.fileno())
cp=json.loads((root/'data/agent-checkpoints.json').read_text());assert '1547581942474608720' in json.dumps(cp)
print(json.dumps({'verification':verification,'inventory_readback':ident,'file_bytes':output.stat().st_size,'skill_readback': 'HTTP201' in paths[-1].read_text()},ensure_ascii=False))
