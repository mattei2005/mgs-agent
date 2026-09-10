import json,hashlib
from pathlib import Path
from collections import defaultdict,Counter
from decimal import Decimal as D
from openpyxl import load_workbook
P=Path(__file__).parent;OLD=Path('/root/mgs-agent/work/gam-sep01-09-1547581942474608720')
files={'original':Path('/root/.hermes/profiles/zeus/cache/documents/doc_10d5b4efd431_report_1_de_set_ate_9_de_set.xlsx'),'zeus':OLD/'Receita-01-09set-2026-Zeus.xlsx','claude':Path('/root/.hermes/profiles/zeus/cache/documents/doc_a73172281346_Receita-01-09set-2026-claude.xlsx')}
manifest={k:{'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for k,p in files.items()}
original=defaultdict(lambda:D(0));original_site=defaultdict(lambda:D(0));n=0
for s in load_workbook(files['original'],read_only=True,data_only=True):
 for row in s.iter_rows(min_row=2,values_only=True):
  if not any(x is not None for x in row):continue
  original[(s.title.upper(),row[0].date().isoformat())]+=D(str(row[9]));n+=1
assert n==27091
books={};labels={};metadata={}
for name in ('zeus','claude'):
 data={};label={};meta={}
 w=load_workbook(files[name],read_only=True,data_only=True)
 for currency in ('USD','CAD'):
  s=w['Receita '+currency];assert list(next(s.values))==['Data','Site','Vertical','Gestor','Receita']
  rows=0
  for i,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
   if row[0]=='TOTAL':label[currency]=D(str(row[4]));continue
   assert len(row)==5 and isinstance(row[4],(float,int))
   key=(currency,*row[:4]);assert key not in data,(name,key);data[key]=D(str(row[4]));rows+=1
  meta[currency]=rows
 books[name]=data;labels[name]=label;metadata[name]=meta
# Baseline comparison; current owner decision is a separate overlay, never rewrite sent baseline.
updated=defaultdict(lambda:D(0))
for k,v in books['zeus'].items():
 nk=k[:4]+('g002-d',) if k[2] in ('escalatepower.com','mavroa.com') else k
 updated[nk]+=v
books['zeus_owner_confirmed']=dict(updated)
raw_agg={tuple(r['key']):D(r['revenue']) for r in json.loads((OLD/'aggregate.json').read_text())}
lineage=json.loads((OLD/'lineage.json').read_text());assert len(lineage)==n
currdays=[]
for k,raw in sorted(original.items()):
 a=sum((v for key,v in books['zeus'].items() if key[:2]==k),D(0));b=sum((v for key,v in books['claude'].items() if key[:2]==k),D(0))
 currdays.append({'currency':k[0],'day':k[1],'original':str(raw),'zeus':str(a),'claude':str(b),'claude_minus_original':str(b-raw),'zeus_minus_original':str(a-raw),'cent_match':a.quantize(D('.01'))==b.quantize(D('.01'))==raw.quantize(D('.01'))})
comparisons={}
for name in ('zeus','zeus_owner_confirmed'):
 a=books[name];b=books['claude'];diff=[];equal=[]
 for k in sorted(set(a)|set(b)):
  av=a.get(k,D(0));bv=b.get(k,D(0));delta=bv-av
  if k in a and k in b and abs(delta)<=D('.00000051'):equal.append(k)
  else:diff.append({'key':k,'zeus':str(av),'claude':str(bv),'claude_minus_zeus':str(delta),'zeus_present':k in a,'claude_present':k in b})
 comparisons[name]={'matching_groups_at_6_decimals':len(equal),'differing_union_keys':len(diff),'differences':diff,'only_zeus':sum(k not in b for k in a),'only_claude':sum(k not in a for k in b)}
(P/'all-differences.json').write_text(json.dumps(comparisons,ensure_ascii=False,indent=2))
# Project onto site/vertical/manager to describe business differences without daily verbosity.
projections={}
for dims,title in [((0,2),'site'),((0,2,3),'site_vertical'),((0,2,4),'site_manager')]:
 ab=[]
 for name in ('zeus_owner_confirmed','claude'):
  dest=defaultdict(lambda:D(0))
  for k,v in books[name].items():dest[tuple(k[i] for i in dims)]+=v
  ab.append(dest)
 a,b=ab;projections[title]=[{'key':k,'zeus':str(a[k]),'claude':str(b[k]),'difference':str(b[k]-a[k])} for k in sorted(set(a)|set(b)) if abs(b[k]-a[k])>D('.00001')]
# Explain raw-source characteristics of every site with a differing key.
changed_sites=sorted({r['key'][2] for r in comparisons['zeus_owner_confirmed']['differences']})
source_details=[]
for site in changed_sites:
 rows=[r for r in lineage if r['key'][2]==site];by=defaultdict(lambda:D(0))
 for r in rows:by[(r['placement'],r['medium'])]+=D(r['amount'])
 source_details.append({'site':site,'source_rows':len(rows),'by_placement_medium':[{'placement':k[0],'medium':k[1],'revenue':str(v)} for k,v in sorted(by.items())]})
(P/'source-details.json').write_text(json.dumps(source_details,ensure_ascii=False,indent=2))
summary={'manifest':manifest,'original_rows':n,'groups':metadata,'totals':{c:{'original':str(sum((v for k,v in original.items() if k[0]==c),D(0))),'zeus_data_sum':str(sum((v for k,v in books['zeus'].items() if k[0]==c),D(0))),'claude_data_sum':str(sum((v for k,v in books['claude'].items() if k[0]==c),D(0))),'zeus_total_label':str(labels['zeus'][c]),'claude_total_label':str(labels['claude'][c])} for c in ('USD','CAD')},'all_currency_day_cent_match':all(r['cent_match'] for r in currdays),'max_claude_daily_raw_difference':str(max(abs(D(r['claude_minus_original'])) for r in currdays)),'comparison_counts':{name:{k:v for k,v in res.items() if k!='differences'} for name,res in comparisons.items()},'changed_sites':changed_sites,'projection_differences':projections,'currency_days':currdays}
(P/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in summary.items() if k not in ('manifest','currency_days')},ensure_ascii=False,indent=2))
