import json,hashlib,shutil,zipfile
from pathlib import Path
from decimal import Decimal as D,getcontext
from collections import defaultdict,Counter
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
getcontext().prec=50
P=Path(__file__).parent
BASE=Path('/root/mgs-agent/work/gam-sep01-09-1547581942474608720/Receita-01-09set-2026-Zeus.xlsx')
ORIGINAL=Path('/root/.hermes/profiles/zeus/cache/documents/doc_10d5b4efd431_report_1_de_set_ate_9_de_set.xlsx')
CLAUDE=Path('/root/.hermes/profiles/zeus/cache/documents/doc_a73172281346_Receita-01-09set-2026-claude.xlsx')
YOLO=Path('/root/mgs-agent/work/yolo-approved-1547678046553636905/approved-allocation.json')
OUT=P/'Receita-01-09set-2026-Zeus-Final.xlsx'
EXPECTED={'original':'c98122ed765a46c25656a57ccd96e73689ac2249166b18ff3d08902cf9da771d','baseline':'85b511036ae56ecf92bcfe44fb5f0f169a7eaf46fb6bf9db958be141c87786d5','claude':'895f0c92ae62e25787b472af4bc9bce2437a7fca0084779ce037a2e885638113'}
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
assert {'original':sha(ORIGINAL),'baseline':sha(BASE),'claude':sha(CLAUDE)}==EXPECTED
# Read exact source controls.
raw_daily=defaultdict(D);raw_count=Counter();ow=load_workbook(ORIGINAL,read_only=True,data_only=True)
assert ow.sheetnames==['usd','cad']
for s in ow:
 c=s.title.upper()
 for r in s.iter_rows(min_row=2,values_only=True):
  if all(v is None for v in r):continue
  day=r[0].strftime('%Y-%m-%d');raw_daily[(c,day)]+=D(str(r[9]));raw_count[c]+=1
ow.close();assert raw_count==Counter({'CAD':25374,'USD':1717})
# Read immutable Zeus baseline aggregates.
bw=load_workbook(BASE,read_only=True,data_only=True);base={};base_daily_site=defaultdict(D)
for c in ['USD','CAD']:
 for r in bw['Receita '+c].iter_rows(min_row=2,values_only=True):
  if r[0]=='TOTAL':continue
  k=(c,*r[:4]);assert k not in base;base[k]=D(str(r[4]));base_daily_site[(c,r[0],r[1])]+=D(str(r[4]))
bw.close()
# Apply only confirmed overlays; independent baselines remain untouched.
final=defaultdict(D);changes=[]
for k,v in base.items():
 c,day,site,vertical,manager=k
 if site=='yolokfx.com':continue
 nv,nm=vertical,manager
 if site in {'cliquet.com','finance.topfeed.fun','openzed.com'} and vertical=='br-cc-br':nv='br-car-br'
 if site=='escalatepower.com':nv,nm='us-cc-en','g002-d'
 if site=='mavroa.com':nv,nm='us-shein-es','g002-d'
 nk=(c,day,site,nv,nm);final[nk]+=v
 if nk!=k:changes.append({'from':k,'to':nk,'amount':str(v)})
y=json.loads(YOLO.read_text());assert y['authorization']=='1547678046553636905' and y['all_9_days_reconciled']
for r in y['rows']:
 v=D(r['revenue'])
 if v==0:continue
 k=(r['currency'],r['date'],r['site'],r['vertical'],r['manager']);final[k]+=v
# Validate overlay scope and every daily/site control.
final_daily_site=defaultdict(D)
for k,v in final.items():final_daily_site[k[:3]]+=v
assert set(final_daily_site)==set(base_daily_site)
assert max(abs(final_daily_site[k]-base_daily_site[k]) for k in base_daily_site)<D('0.000000001')
assert not any('A confirmar' in k for k in final)
assert not any(k[2]=='mavroa.com' and (k[3],k[4])!=('us-shein-es','g002-d') for k in final)
assert not any(k[2]=='escalatepower.com' and (k[3],k[4])!=('us-cc-en','g002-d') for k in final)
for site in ['cliquet.com','finance.topfeed.fun','openzed.com']:
 assert not any(k[2]==site and k[3]=='br-cc-br' for k in final)
# Build the reviewed workbook using baseline theme.
book=load_workbook(BASE)
navy='17365D';blue='DCE6F1';green='E2F0D9';amber='FFF2CC';thin='D9E2F3'
for c in ['USD','CAD']:
 s=book['Receita '+c]
 s.delete_rows(2,s.max_row-1)
 rows=[(k,v) for k,v in sorted(final.items()) if k[0]==c]
 for k,v in rows:
  s.append([*k[1:],float(v)])
  for cell in s[s.max_row]:cell.alignment=Alignment(vertical='top')
  if k[2]=='yolokfx.com':
   s.cell(s.max_row,5).comment=Comment('Total diário controlado pelo GAM. G001/G003/G004/G005/G006 usam receita bruta SB por conta; G002-s recebe o complemento diário. Decisão 1547678046553636905.','Zeus')
  if (k[2],k[3]) in {('cliquet.com','br-car-br'),('finance.topfeed.fun','br-car-br'),('openzed.com','br-car-br')}:
   s.cell(s.max_row,3).comment=Comment('Vertical BR confirmada por Rodolfo nesta revisão; país e receita preservados do GAM.','Zeus')
  if k[2]=='mavroa.com':s.cell(s.max_row,3).comment=Comment('us-shein-es confirmado por Rodolfo em 1547689598392737874.','Zeus')
  if k[2]=='escalatepower.com':s.cell(s.max_row,3).comment=Comment('us-cc-en confirmado por Rodolfo em 1547688086664908952.','Zeus')
 total=sum((v for k,v in final.items() if k[0]==c),D(0));s.append(['TOTAL',None,None,None,float(total)])
 for cell in s[s.max_row]:cell.fill=PatternFill('solid',fgColor=blue);cell.font=Font(bold=True,color=navy)
 for row in s.iter_rows(min_row=2):row[4].number_format='#,##0.00'
 s.auto_filter.ref=f'A1:E{s.max_row-1}';s.freeze_panes='A2'
# Rebuild concise validation sheet.
book.remove(book['Validacao']);v=book.create_sheet('Validacao')
v.append(['Controle / data','Moeda / site','Original GAM / outros gestores','Excel final / G002','Diferença / observação'])
for cell in v[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor=navy)
v.append(['Período','01 a 09/09/2026',None,None,'Datas lidas do original GAM recebido por e-mail.'])
v.append(['Fonte principal',ORIGINAL.name,None,None,'Receita bruta GAM; sem câmbio, share, imposto ou substituição por total da dashboard.'])
v.append(['Linhas originais','USD',raw_count['USD'],None,None]);v.append(['Linhas originais','CAD',raw_count['CAD'],None,None])
for c in ['USD','CAD']:
 for day in [f'2026-09-{d:02}' for d in range(1,10)]:
  got=sum((x for k,x in final.items() if k[:2]==(c,day)),D(0));orig=raw_daily[(c,day)]
  v.append([day,c,float(orig),float(got),float(got-orig)])
 total=sum((x for k,x in raw_daily.items() if k[0]==c),D(0));got=sum((x for k,x in final.items() if k[0]==c),D(0))
 v.append(['TOTAL '+c,c,float(total),float(got),float(got-total)])
v.append([]);v.append(['YOLOKFX','GAM diário = outros gestores SB + G002-s',None,None,'Decisão 1547678046553636905; não extrapolar ao próximo relatório.'])
for d in y['daily']:v.append([d['date'],'yolokfx.com',float(D(d['sb_other_managers'])),float(D(d['g002_residual'])),f"Total GAM CAD {D(d['gam']):.6f}; fechado por dia."])
v.append([]);v.append(['DECISÕES APLICADAS','Escopo 01–09/09/2026',None,None,'Baselines Zeus/Claude preservados por hash.'])
decisions=[
('Openzed Finanças','es-cc-es e us-cc-es separados','1547598180990976100'),('Eggbev','gb-cc-en e us-cc-en separados; g006-d','1547674720982278194'),('Yolokfx','outros gestores SB; G002-s complemento GAM diário','1547678046553636905'),('Autocreditadx','us-car-en; g002-d','1547679210892566578'),('Cliquet','GB/US separados; BR br-car-br; g002-d','1547681384653529169'),('Zytiva Finanças','es-cc-es e us-cc-es separados; g003-d','1547682605766414356'),('Topfeed','BR br-car-br; g004-d; US/GB preservados','1547683789034102797'),('Openzed','BR br-car-br; g003-d; US/GB preservados','1547684367349059627'),('Escalatepower','us-cc-en; g002-d','1547688086664908952'),('Mavroa','us-shein-es; g002-d','1547689598392737874')]
for site,rule,mid in decisions:v.append([site,rule,None,None,'Rodolfo '+mid])
v.append([]);v.append(['REGRA FUTURA','Preservar país que vier no próximo GAM',None,None,'Decisões desta revisão não remapeiam automaticamente novos países/períodos.'])
v.append(['PENDÊNCIAS','Nenhuma classificação pendente neste arquivo',0,0,'Todos os domínios divergentes foram decididos por Rodolfo.'])
v.freeze_panes='A2';v.sheet_view.showGridLines=False
for col,width in {'A':34,'B':52,'C':28,'D':27,'E':92}.items():v.column_dimensions[col].width=width
for row in v.iter_rows(min_row=2):
 for cell in row:
  if isinstance(cell.value,(float,int)):cell.number_format='#,##0.000000'
  cell.alignment=Alignment(vertical='top',wrap_text=True)
 v.row_dimensions[row[0].row].height=32
book.calculation.fullCalcOnLoad=True;book.calculation.forceFullCalc=True
book.save(OUT)
# Independent serialized readback.
rw=load_workbook(OUT,read_only=False,data_only=False);assert rw.sheetnames==['Receita USD','Receita CAD','Validacao'];check={};totals={};formulas=[]
for c in ['USD','CAD']:
 s=rw['Receita '+c];assert s.freeze_panes=='A2' and s.auto_filter.ref==f'A1:E{s.max_row-1}'
 for r in s.iter_rows(min_row=2,values_only=True):
  if r[0]=='TOTAL':totals[c]=D(str(r[4]));continue
  k=(c,*r[:4]);assert k not in check;check[k]=D(str(r[4]))
for s in rw:
 for row in s.iter_rows():
  for cell in row:
   if isinstance(cell.value,str) and cell.value.startswith('='):formulas.append((s.title,cell.coordinate,cell.value))
rw.close();assert not formulas and set(check)==set(final)
max_group_error=max(abs(check[k]-final[k]) for k in final)
assert max_group_error<D('0.000000001')
serial_daily={}
for c,day in raw_daily:
 got=sum((x for k,x in check.items() if k[:2]==(c,day)),D(0));serial_daily[(c,day)]=got;assert abs(got-raw_daily[(c,day)])<D('0.00000001')
for c in ['USD','CAD']:
 target=sum((x for k,x in raw_daily.items() if k[0]==c),D(0));assert abs(totals[c]-target)<D('0.00000001')
with zipfile.ZipFile(OUT) as z:bad=z.testzip();assert bad is None
# Baselines must still be byte-identical.
assert {'original':sha(ORIGINAL),'baseline':sha(BASE),'claude':sha(CLAUDE)}==EXPECTED
summary={'status':'FINAL_REVIEWED_RECONCILED','authorization':'1547689598392737874','input':str(ORIGINAL),'baseline':str(BASE),'claude':str(CLAUDE),'output':str(OUT),'hashes':{'original':sha(ORIGINAL),'baseline':sha(BASE),'claude':sha(CLAUDE),'output':sha(OUT)},'input_rows':dict(raw_count),'output_groups':len(final),'groups_by_currency':dict(Counter(k[0] for k in final)),'raw_totals':{c:str(sum((x for k,x in raw_daily.items() if k[0]==c),D(0))) for c in ['USD','CAD']},'serialized_totals':{c:str(totals[c]) for c in ['USD','CAD']},'max_serialized_group_error':str(max_group_error),'all_18_currency_days_verified':True,'all_daily_site_totals_preserved':True,'no_pending_classifications':True,'no_formulas_or_formula_errors':True,'xlsx_zip_integrity':True,'baseline_hashes_unchanged':True,'financial_imports':False,'confirmed_changes':changes,'yolo_totals':y['totals']}
normalized=json.loads(json.dumps(summary,ensure_ascii=False))
(P/'summary.json').write_text(json.dumps(normalized,ensure_ascii=False,indent=2)+'\n')
(P/'final-aggregate.json').write_text(json.dumps([{'key':k,'revenue':str(x)} for k,x in sorted(final.items())],ensure_ascii=False,indent=2)+'\n')
assert json.loads((P/'summary.json').read_text())==normalized
print(json.dumps({'status':summary['status'],'output':str(OUT),'sha256':summary['hashes']['output'],'groups':summary['output_groups'],'groups_by_currency':summary['groups_by_currency'],'raw_totals':summary['raw_totals'],'max_group_error':summary['max_serialized_group_error'],'all_18_days':True,'pending':0,'baseline_hashes_unchanged':True},ensure_ascii=False))
