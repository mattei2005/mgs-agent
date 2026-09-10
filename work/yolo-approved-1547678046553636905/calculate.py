import json,re,hashlib
from pathlib import Path
from collections import defaultdict
from decimal import Decimal as D,getcontext
from openpyxl import load_workbook
getcontext().prec=50
P=Path(__file__).parent
B=Path('/root/mgs-agent/work/gam-sep01-09-1547581942474608720')
s=json.loads((B/'summary.json').read_text());src=Path(s['input'])
assert hashlib.sha256(src.read_bytes()).hexdigest()==s['source_sha256']
gam=defaultdict(D);raw_count=0
w=load_workbook(src,read_only=True,data_only=True)
for r in w['cad'].iter_rows(min_row=2,values_only=True):
 if r[1]=='pl_digital-trust_yolokfx_us':
  gam[r[0].strftime('%Y-%m-%d')]+=D(str(r[9]));raw_count+=1
w.close()
assert len(gam)==9 and raw_count==340
batches=json.loads((B/'sb-daily.json').read_text());sb=defaultdict(lambda:defaultdict(D));seen=set()
for b in batches:
 assert b['request']['currency']=='CAD' and b['request']['publishers']==['digital-trust_yolokfx']
 for r in b['data']:
  assert r['DOMAIN']=='yolokfx' and r['DATE'] in gam
  pk=r['PK_JBF_PERFORMANCE_PER_ADGROUP'];assert pk not in seen;seen.add(pk)
  m=re.search(r'(?:^|[-\s])G(00[1-6])(?:\s|$|[-(])',r.get('ACCOUNT_NAME') or '',re.I)
  manager='g'+m.group(1)+'-s' if m else 'unidentified'
  sb[r['DATE']][manager]+=D(str(r['REVENUE'] or 0))
assert len(seen)==288
rows=[];daily=[];totals=defaultdict(D)
others=['g001-s','g003-s','g004-s','g005-s','g006-s']
for day in sorted(gam):
 alloc={m:sb[day][m] for m in others}
 alloc['g002-s']=gam[day]-sum(alloc.values())
 assert alloc['g002-s']>=0 and sum(alloc.values())==gam[day]
 daily.append({'date':day,'gam':str(gam[day]),'sb_other_managers':str(sum(sb[day][m] for m in others)),'g002_residual':str(alloc['g002-s']),'reconciled':True})
 for m,v in sorted(alloc.items()):
  totals[m]+=v;rows.append({'currency':'CAD','date':day,'site':'yolokfx.com','vertical':'us-shein-en','manager':m,'revenue':str(v),'provenance':'GAM_daily_total_minus_other_SB_managers' if m=='g002-s' else 'SB_AdGroup_gross_REVENUE_by_ACCOUNT_NAME'})
assert sum(totals.values())==sum(gam.values())
out={'authorization':'1547678046553636905','supersedes':'Yolo GAM campaign-identity bridge allocation for September1-9 only','source_sha256':s['source_sha256'],'sb_sha256':hashlib.sha256((B/'sb-daily.json').read_bytes()).hexdigest(),'original_rows':raw_count,'unique_sb_rows':len(seen),'daily':daily,'rows':rows,'totals':{k:str(v) for k,v in sorted(totals.items())},'gam_total':str(sum(gam.values())),'all_9_days_reconciled':True,'baseline_workbooks_unchanged':True,'financial_imports':False}
p=P/'approved-allocation.json';p.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n');assert json.loads(p.read_text())==out
print(json.dumps({'status':'PASS','totals':{k:format(v,'.2f') for k,v in sorted(totals.items())},'total':format(sum(gam.values()),'.2f'),'daily_reconciled':len(daily),'negative_residuals':False,'path':str(p)}))
