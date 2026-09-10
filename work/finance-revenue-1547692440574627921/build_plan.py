import json,re,hashlib
from pathlib import Path
from collections import defaultdict,Counter
from decimal import Decimal as D,getcontext
from openpyxl import load_workbook
getcontext().prec=50
P=Path(__file__).parent
X=Path('/root/mgs-agent/work/gam-final-1547689598392737874/Receita-01-09set-2026-Zeus-Final.xlsx')
W=P/'workspace-before.json';OUT=P/'import-plan.json'
EXPECTED='f7047d7ab39a1412a969bc3528bc9b5c6fd763306ad2145ee98a545c81c55477';assert hashlib.sha256(X.read_bytes()).hexdigest()==EXPECTED
w=json.loads(W.read_text());assert w['id']=='workspace-2026-09' and w['state']=='draft'
# Dashboard's verified domain aliases, plus three CSV-only sites kept as pending native facts.
domains={"eggbev.com":"Eggbev","finanzas.eggbev.com":"Eggbev Finanzas","lyzmo.com":"Lyzmo","finanzas.lyzmo.com":"Lyzmo Finanzas","finance.topfeed.fun":"FinanceTopFeed","finanzas.topfeed.fun":"TopFeed Finanzas","finanzas.zuout.com":"Zuout Finanzas","zytiva.com":"Zytiva","finanzas.zytiva.com":"Zytiva Finanzas","newsoun.com":"Newsoun","finanzas.newsoun.com":"Newsoun Finanzas","de.newsoun.com":"Newsoun DE","openzed.com":"Openzed","finanzas.openzed.com":"Openzed Finanzas","cliquet.com":"Cliquet","finanzas.cliquet.com":"Cliquet Finanzas","fincgriffin.com":"Fincgriffin","financeadx.com":"FinanceAdx","marevelx.com":"Marevelx","helixenit.com":"Helixenit","infinitynexx.com":"Infinitynexx","vizioid.com":"Vizioid","xyvlov.com":"Xyvlov","wavesbee.com":"WavesBee","portalrelevante.com":"Portal Relevante","gamingadx.com":"GamingAdx","gamezonead.com":"GameZoneAd","creditoparaveiculo.com":"CreditoParaVeiculo","autocreditadx.com":"AutoCreditAdx","cephyric.com":"Cephyric","escalatepower.com":"Escalatepower","mavroa.com":"Mavroa","yolokfx.com":"Yolokfx"}
# Missing country/site combinations are native facts so CSV geography is never remapped to a dashboard default.
manager_map={'g001':'george','g002':'SEM_COMISSAO','g003':'isliago','g004':'joe','g005':'kelly','g006':'nicolas'}
facts={f['id']:f for f in w['domain']['facts']};model=w['model'];known_sites={f['site'] for f in facts.values()};site_cfg={s['name']:s for s in w['sites']}
rates={r['source']:str(r['value']) for r in w['rates']};quotes={'USDBRL':str(w['fx']),'USDCAD':rates['H1'],'GBPUSD':rates['I1']}
regular=defaultdict(D);native=defaultdict(D);lineage=[];source_totals=defaultdict(D);source_rows=0;collisions=Counter()
book=load_workbook(X,read_only=True,data_only=True)
for currency in ['USD','CAD']:
 for r in book['Receita '+currency].iter_rows(min_row=2,values_only=True):
  if r[0]=='TOTAL':continue
  date,domain,vertical,tag,value=r;value=D(str(value));source_rows+=1;source_totals[(currency,date)]+=value
  assert domain in domains and re.fullmatch(r'[a-z]{2}-[a-z]+-[a-z]{2}',vertical), (domain,vertical)
  site=domains[domain];country=vertical.split('-')[0].upper();m=re.fullmatch(r'(g00[1-6])-[ds]',tag);assert m,(domain,tag);manager=manager_map[m.group(1)]
  candidates=[]
  # CAD stays in its original currency through native facts; legacy dashboard gross cells are USD-base fields.
  if currency=='USD' and domain!='yolokfx.com':
   for fid,f in facts.items():
    if (f['site'],f['country'],f['date'])!=(site,country,date):continue
    for key in model['facts'].get(fid,{}).get('gross',[]):
     x=model['inputs'][key];origin='USD' if x.get('book')=='principal' and x.get('source')==f.get('source',{}).get('gross') else x.get('currency')
     ims=x.get('managers') or []
     identity=(manager in ims) if manager!='SEM_COMISSAO' else not ims
     if origin==currency and identity:candidates.append((key,fid,ims))
  has_country=any((f['site'],f['country'])==(site,country) for f in facts.values())
  if len(candidates)==1:
   key,fid,ims=candidates[0];regular[key]+=value;collisions[key]+=1;lineage.append({'date':date,'domain':domain,'vertical':vertical,'manager':tag,'currency':currency,'revenue':str(value),'route':'dashboard_input','key':key,'fact_id':fid,'dashboard_site':site})
  else:
   assert not candidates and (currency=='CAD' or domain=='yolokfx.com' or not has_country),(domain,date,tag,currency,'country_present='+str(has_country),candidates)
   reason='Yolokfx manager split' if domain=='yolokfx.com' else 'Original CAD preserved through currency-aware native fact' if currency=='CAD' else 'Country/site absent in dashboard; CSV country preserved'
   route=('native',date,site,country,manager,currency,vertical,tag);native[route]+=value;lineage.append({'date':date,'domain':domain,'vertical':vertical,'manager':tag,'currency':currency,'revenue':str(value),'route':'native_manager_fact','dashboard_site':site,'reason':reason})
book.close();assert source_rows==513
# Target inputs must be untouched before the authorized import.
for key in regular:assert str(model['inputs'][key].get('value',''))=='',(key,model['inputs'][key].get('value'))
entries=[];prefix='gam-2026-09-01-09-f7047d7ab39a'
for (_,date,site,country,manager,currency,vertical,tag),value in sorted(native.items()):
 slug=re.sub(r'[^a-z0-9]+','-',site.lower()).strip('-');id=f'{prefix}|{slug}|{country}|{date[-2:]}|{tag}|{currency}'
 cfg=site_cfg.get(site);partner=(cfg or {}).get('network') or 'SB Rede1'
 entries.append({'id':id,'source_import_id':prefix,'source_sha256':EXPECTED,'source_vertical':vertical,'source_manager_tag':tag,'site':site,'partner':partner,'manager':manager,'country':country,'date':date,'currency':currency,'gross':str(value),'spend':'0','invalid_rate':rates['L1'],'share_rate':rates['D1'],'tax_rate':rates['C1'],'quotes':quotes})
assert not any((a.get('source_import_id')==prefix or str(a.get('id','')).startswith(prefix)) for a in w['additions'])
changes=[{'key':k,'value':str(v)} for k,v in sorted(regular.items())]
# Full financial and routing reconciliation.
plan_totals=defaultdict(D)
for x in changes:
 inp=model['inputs'][x['key']];fid=next(fid for fid,f in model['facts'].items() if x['key'] in f.get('gross',[]));f=facts[fid];origin='USD' if inp.get('book')=='principal' and inp.get('source')==f.get('source',{}).get('gross') else inp.get('currency');plan_totals[(origin,f['date'])]+=D(x['value'])
for x in entries:plan_totals[(x['currency'],x['date'])]+=D(x['gross'])
assert set(plan_totals)==set(source_totals) and max(abs(plan_totals[k]-source_totals[k]) for k in source_totals)<D('0.000000001')
coverage=Counter(x['route'] for x in lineage);assert sum(coverage.values())==513
native_sites=sorted({x['site'] for x in entries});assert {x['domain'] for x in lineage if x['route']=='dashboard_input'}=={'fincgriffin.com'}
assert {x['domain'] for x in lineage if x['currency']=='CAD'}<={x['domain'] for x in lineage if x['route']=='native_manager_fact'}
plan={'authorization_message_id':'1547692440574627921','period':'2026-09','date_range':['2026-09-01','2026-09-09'],'scenario_id':w['id'],'preflight_revision':w['revision'],'source_excel':str(X),'source_sha256':EXPECTED,'source_rows':source_rows,'source_totals':{f'{c}|{d}':str(v) for (c,d),v in sorted(source_totals.items())},'changes':changes,'native_entries':entries,'lineage':lineage,'summary':{'mapped_excel_rows':coverage['dashboard_input'],'native_excel_rows':coverage['native_manager_fact'],'input_fields':len(changes),'input_field_collisions':sum(n-1 for n in collisions.values()),'native_entries':len(entries),'native_sites':native_sites,'all_18_currency_days_reconciled':True,'existing_target_gross_fields_nonempty':0,'existing_import_rows':0,'strategy_suffix_preserved_in_excel_only':True,'dashboard_manager_identity_map':manager_map,'unknown_sites_registered':False,'company_expense_allocation_changed':False,'financial_imports_before_apply':0}}
OUT.write_text(json.dumps(plan,ensure_ascii=False,indent=2)+'\n');assert json.loads(OUT.read_text())==plan
print(json.dumps({'status':'PASS',**plan['summary'],'source_rows':source_rows,'output':str(OUT)},ensure_ascii=False))
