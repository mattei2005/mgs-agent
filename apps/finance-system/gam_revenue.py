"""Fail-closed parser/planner for daily emailed GAM revenue reports."""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal, getcontext
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

getcontext().prec = 50
ROOT = Path(__file__).resolve().parent
RULES_PATH = Path("/root/mgs-agent/data/finance-gam-revenue-rules.json")
VALID_MANAGER = re.compile(r"^(g00[1-6])-(d|s)$")
PLACEMENT = re.compile(r"^pl_digital-trust_(.+)_([a-z]{2})$")
VERTICAL = re.compile(r"^[a-z]{2}-[a-z]+-[a-z]{2}$")

REPORTS = {
    "usd": {
        "currency": "USD",
        "report_name": "Report Digital Trust (adx 2)",
        "network": "All Digital Marketing",
        "attachment": "Report Digital Trust (adx 2).xlsx",
    },
    "cad": {
        "currency": "CAD",
        "report_name": "Digital Trust",
        "network": "00-JBF Digital Server",
        "attachment": "Digital Trust.xlsx",
    },
}


def norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).lower().split())


def load_rules(path: Path = RULES_PATH) -> dict[str, Any]:
    data = json.loads(path.read_text())
    assert data["schema_version"] == 1
    return data


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _property_map(ws) -> dict[str, str]:
    out: dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        if len(row) >= 2 and row[0] is not None:
            out[norm(row[0])] = str(row[1] or "").strip()
    return out


def _prop(props: dict[str, str], *labels: str) -> str:
    for label in labels:
        if norm(label) in props:
            return props[norm(label)]
    return ""


def _header_index(headers: list[str], *terms: str) -> int:
    wanted = [norm(x) for x in terms]
    for index, header in enumerate(headers):
        if all(term in norm(header) for term in wanted):
            return index
    raise ValueError("missing header: " + "+".join(terms))


def _date(value: Any) -> str:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = str(value or "").strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    raise ValueError("invalid source date")


def inspect_workbook(path: Path, report_key: str) -> dict[str, Any]:
    expected = REPORTS[report_key]
    if path.name != expected["attachment"]:
        raise ValueError(f"unexpected attachment for {report_key}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 2:
            raise ValueError("workbook must contain exactly properties and data sheets")
        props = _property_map(workbook.worksheets[0])
        currency = _prop(props, "Report currency", "Moeda do relatório")
        report_name = _prop(props, "Report name", "Nome do relatório")
        network = _prop(props, "Publisher network", "Rede do publisher")
        timezone = _prop(props, "Time zone", "Fuso horário")
        date_range = _prop(props, "Date range", "Período")
        if currency != expected["currency"]:
            raise ValueError(f"currency mismatch for {report_key}")
        if report_name != expected["report_name"]:
            raise ValueError(f"report-name mismatch for {report_key}")
        if network != expected["network"]:
            raise ValueError(f"network mismatch for {report_key}")
        if timezone != "America/Sao_Paulo":
            raise ValueError(f"timezone mismatch for {report_key}")

        ws = workbook.worksheets[1]
        headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        date_i = _header_index(headers, "date") if any("date" in norm(x) for x in headers) else _header_index(headers, "data")
        placement_i = _header_index(headers, "placement") if any(norm(x) == "placement" for x in headers) else _header_index(headers, "posicao")
        medium_i = _header_index(headers, "utm_medium")
        campaign_i = _header_index(headers, "utm_campaign")
        content_i = _header_index(headers, "utm_content")
        revenue_i = _header_index(headers, "revenue") if any("revenue" in norm(x) for x in headers) else _header_index(headers, "receita")
        rows = []
        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not any(value is not None for value in row):
                continue
            revenue = Decimal(str(row[revenue_i] or 0))
            rows.append(
                {
                    "source_row": row_number,
                    "date": _date(row[date_i]),
                    "placement": str(row[placement_i] or "").strip(),
                    "medium": str(row[medium_i] or "").strip(),
                    "campaign": str(row[campaign_i] or "").strip(),
                    "content": str(row[content_i] or "").strip(),
                    "revenue": revenue,
                }
            )
        if not rows:
            raise ValueError("empty GAM data sheet")
        dates = sorted({row["date"] for row in rows})
        if len(dates) != 1:
            raise ValueError("daily report contains multiple dates")
        return {
            "report_key": report_key,
            "currency": currency,
            "network": network,
            "timezone": timezone,
            "date_range_property": date_range,
            "date": dates[0],
            "headers": headers,
            "rows": rows,
            "source_sha256": sha256(path),
            "source_path": str(path),
        }
    finally:
        workbook.close()


def _blocker_key(item: dict[str, Any]) -> tuple[str, ...]:
    return tuple(str(item.get(key, "")) for key in ("type", "domain", "country", "date", "medium"))


def build_plan(paths: dict[str, Path], *, rules_path: Path = RULES_PATH) -> dict[str, Any]:
    if set(paths) != set(REPORTS):
        raise ValueError("complete USD+CAD report pair required")
    rules = load_rules(rules_path)
    reports = [inspect_workbook(paths[key], key) for key in sorted(paths)]
    dates = {report["date"] for report in reports}
    if len(dates) != 1:
        raise ValueError("report pair date mismatch")
    source_date = next(iter(dates))
    period = source_date[:7]
    source_hashes = {report["report_key"]: report["source_sha256"] for report in reports}
    bundle_material = json.dumps(source_hashes, sort_keys=True, separators=(",", ":")).encode()
    bundle_hash = hashlib.sha256(bundle_material).hexdigest()

    blockers: dict[tuple[str, ...], dict[str, Any]] = {}
    mapped_rows: list[dict[str, Any]] = []
    source_totals: defaultdict[str, Decimal] = defaultdict(Decimal)
    grouped: defaultdict[tuple[str, str, str, str, str], Decimal] = defaultdict(Decimal)
    fallback_rows = Counter()
    forced_rows = Counter()

    for report in reports:
        currency = report["currency"]
        for row in report["rows"]:
            source_totals[currency] += row["revenue"]
            match = PLACEMENT.fullmatch(row["placement"])
            if not match:
                item = {"type": "unknown_placement", "date": source_date, "placement": row["placement"], "currency": currency, "rows": 0, "revenue": Decimal(0)}
                key = _blocker_key(item) + (row["placement"],)
                target = blockers.setdefault(key, item)
                target["rows"] += 1
                target["revenue"] += row["revenue"]
                continue
            brand, country = match.groups()
            domain = rules["brand_domains"].get(brand)
            if not domain:
                item = {"type": "unknown_domain", "date": source_date, "domain": brand, "country": country, "currency": currency, "rows": 0, "revenue": Decimal(0)}
                target = blockers.setdefault(_blocker_key(item), item)
                target["rows"] += 1
                target["revenue"] += row["revenue"]
                continue
            site = rules["dashboard_sites"].get(domain)
            vertical = rules["vertical_by_domain_country"].get(f"{domain}|{country}")
            if not site or not vertical or not VERTICAL.fullmatch(vertical):
                item = {"type": "new_domain_country", "date": source_date, "domain": domain, "country": country, "currency": currency, "rows": 0, "revenue": Decimal(0)}
                target = blockers.setdefault(_blocker_key(item), item)
                target["rows"] += 1
                target["revenue"] += row["revenue"]
                continue

            original_medium = row["medium"]
            forced = rules.get("force_manager_tag", {}).get(domain)
            manager_tag = forced or (original_medium if VALID_MANAGER.fullmatch(original_medium) else None)
            route = "forced" if forced else "source"
            block_from = rules.get("missing_manager_block_from", {}).get(domain)
            if not manager_tag and block_from and source_date >= block_from:
                item = {"type": "missing_manager_after_cutover", "date": source_date, "domain": domain, "country": country, "currency": currency, "medium": original_medium, "rows": 0, "revenue": Decimal(0), "campaigns": set()}
                target = blockers.setdefault(_blocker_key(item), item)
                target["rows"] += 1
                target["revenue"] += row["revenue"]
                target["campaigns"].add(row["campaign"])
                continue
            if not manager_tag:
                manager_tag = rules.get("fallback_manager_tag", {}).get(domain)
                route = "fallback"
            if not manager_tag or not VALID_MANAGER.fullmatch(manager_tag):
                item = {"type": "unknown_manager", "date": source_date, "domain": domain, "country": country, "currency": currency, "medium": original_medium, "rows": 0, "revenue": Decimal(0), "campaigns": set()}
                target = blockers.setdefault(_blocker_key(item), item)
                target["rows"] += 1
                target["revenue"] += row["revenue"]
                target["campaigns"].add(row["campaign"])
                continue
            manager = rules["manager_identity"][manager_tag[:4]]
            key = (currency, site, country.upper(), vertical, manager_tag)
            grouped[key] += row["revenue"]
            if route == "fallback":
                fallback_rows[domain] += 1
            elif route == "forced" and original_medium != manager_tag:
                forced_rows[domain] += 1
            mapped_rows.append(
                {
                    "report": report["report_key"],
                    "source_row": row["source_row"],
                    "date": source_date,
                    "currency": currency,
                    "placement": row["placement"],
                    "domain": domain,
                    "site": site,
                    "country": country.upper(),
                    "vertical": vertical,
                    "source_medium": original_medium,
                    "manager_tag": manager_tag,
                    "manager": manager,
                    "manager_route": route,
                    "campaign": row["campaign"],
                    "content": row["content"],
                    "revenue": str(row["revenue"]),
                }
            )

    serialized_blockers = []
    for item in blockers.values():
        clean = dict(item)
        clean["revenue"] = str(clean["revenue"])
        if isinstance(clean.get("campaigns"), set):
            clean["campaigns"] = sorted(clean["campaigns"])
        serialized_blockers.append(clean)
    serialized_blockers.sort(key=lambda x: (x["type"], x.get("domain", ""), x.get("country", ""), x.get("currency", "")))

    prefix = f"gam-email-{source_date}-{bundle_hash[:12]}"
    entries = []
    for (currency, site, country, vertical, manager_tag), revenue in sorted(grouped.items()):
        slug = re.sub(r"[^a-z0-9]+", "-", site.lower()).strip("-")
        entries.append(
            {
                "id": f"{prefix}|{slug}|{country}|{manager_tag}|{currency}",
                "source_import_type": "gam_email_daily",
                "source_import_id": prefix,
                "source_date": source_date,
                "source_bundle_sha256": bundle_hash,
                "source_hashes": source_hashes,
                "source_vertical": vertical,
                "source_manager_tag": manager_tag,
                "site": site,
                "manager": rules["manager_identity"][manager_tag[:4]],
                "country": country,
                "date": source_date,
                "currency": currency,
                "gross": str(revenue),
                "spend": "0",
            }
        )

    grouped_totals: defaultdict[str, Decimal] = defaultdict(Decimal)
    for entry in entries:
        grouped_totals[entry["currency"]] += Decimal(entry["gross"])
    if not serialized_blockers and dict(grouped_totals) != dict(source_totals):
        raise AssertionError("mapped totals do not reconcile")

    return {
        "schema_version": 1,
        "authorization_message_id": "1547983130038767755",
        "period": period,
        "date": source_date,
        "scenario_id": f"workspace-{period}",
        "source_import_id": prefix,
        "source_bundle_sha256": bundle_hash,
        "source_hashes": source_hashes,
        "source_files": {report["report_key"]: report["source_path"] for report in reports},
        "source_rows": sum(len(report["rows"]) for report in reports),
        "source_totals": {key: str(value) for key, value in sorted(source_totals.items())},
        "entries": entries,
        "lineage": mapped_rows,
        "blockers": serialized_blockers,
        "summary": {
            "mapped_rows": len(mapped_rows),
            "blocked_rows": sum(item["rows"] for item in serialized_blockers),
            "groups": len(entries),
            "fallback_rows": dict(sorted(fallback_rows.items())),
            "forced_rows": dict(sorted(forced_rows.items())),
            "currency_totals_reconciled": not serialized_blockers and dict(grouped_totals) == dict(source_totals),
        },
    }


def write_plan(paths: dict[str, Path], output: Path, *, rules_path: Path = RULES_PATH) -> dict[str, Any]:
    plan = build_plan(paths, rules_path=rules_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.name + ".pending")
    temporary.write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n")
    temporary.chmod(0o600)
    temporary.replace(output)
    return plan
